# -*- coding: utf-8 -*-
"""Genera el PPTX a partir del deck de relanzamiento-dci.vercel.app.
Sistema de marca DCI Brand Evolution (junio 2026).
Canela y Sweet Sans Pro no tienen licencia local: se sustituyen por
Baskerville y Avenir Next, que es lo mas cercano que trae macOS."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn

SC = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(SC), 'out.pptx')

# ---------------------------------------------------------------- tokens
NAVY = C(0x04, 0x1E, 0x42)
BLUE = C(0x00, 0x4C, 0x97)
SKY  = C(0xCC, 0xEC, 0xF9)
CREAM= C(0xF5, 0xF3, 0xC4)
ICE  = C(0xF3, 0xFA, 0xFA)
SAND = C(0xD7, 0xBC, 0x81)
WHITE= C(0xFF, 0xFF, 0xFF)
MUTED= C(0x5A, 0x6B, 0x85)   # rgba(4,30,66,.64) sobre ice
LINE = C(0xD3, 0xD8, 0xE1)
SKYM = C(0xA8, 0xC4, 0xDA)   # sky al 82% sobre navy
OUTL = C(0x3A, 0x5A, 0x7E)   # borde de .c.ou sobre navy

SERIF = 'Baskerville'
SANS  = 'Avenir Next'

W, H = 13.333, 7.5
M = 0.8                 # margen lateral
CW = W - 2 * M          # 11.733

prs = Presentation()
prs.slide_width  = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

LOGO_L = os.path.join(SC, 'logo_lg-l.png')   # logo oscuro, para fondo claro
LOGO_N = os.path.join(SC, 'logo_lg-n.png')   # logo claro, para fondo navy
EMB_L  = os.path.join(SC, 'logo_emb-l.png')
EMB_N  = os.path.join(SC, 'logo_emb-n.png')


# ---------------------------------------------------------------- helpers
def _spc(run, hundredths):
    """letter-spacing: el atributo spc de a:rPr va en centesimas de punto."""
    run.font._rPr.set('spc', str(int(hundredths)))


def _nospace(tf):
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True


def rect(sl, l, t, w, h, fill=None, line=None, radius=None):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = sl.shapes.add_shape(shp, Inches(l), Inches(t), Inches(w), Inches(h))
    if radius:
        # adj = radio / lado menor
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.fill.solid(); s.line.fill.fore_color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def text(sl, l, t, w, h, parts, font=SANS, size=10, bold=False,
         color=NAVY, space=0, lead=None, spc=0, anchor=None, italic=False):
    """parts: str | [str] | [[(txt, {opts}), ...]]  (lista de parrafos de runs)"""
    s = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = s.text_frame
    _nospace(tf)
    if anchor:
        tf.vertical_anchor = anchor
    if isinstance(parts, str):
        parts = [parts]
    for i, para in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if space:
            p.space_after = Pt(space)
        if lead:
            p.line_spacing = lead
        runs = para if isinstance(para, list) else [(para, {})]
        for txt, o in runs:
            r = p.add_run(); r.text = txt
            r.font.name  = o.get('font', font)
            r.font.size  = Pt(o.get('size', size))
            r.font.bold  = o.get('bold', bold)
            r.font.italic= o.get('italic', italic)
            r.font.color.rgb = o.get('color', color)
            sp = o.get('spc', spc)
            if sp:
                _spc(r, sp)
    return s


def bullets(sl, l, t, w, h, items, color=NAVY, size=9.2, dot=BLUE, space=3.5):
    """items: str o [(bold, resto)]"""
    s = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = s.text_frame; _nospace(tf)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        p.line_spacing = 1.18
        r0 = p.add_run(); r0.text = u'•  '
        r0.font.name = SANS; r0.font.size = Pt(size); r0.font.color.rgb = dot; r0.font.bold = True
        if isinstance(it, tuple):
            b, rest = it
            rb = p.add_run(); rb.text = b
            rb.font.name = SANS; rb.font.size = Pt(size); rb.font.bold = True; rb.font.color.rgb = color
            if rest:
                rr = p.add_run(); rr.text = rest
                rr.font.name = SANS; rr.font.size = Pt(size); rr.font.color.rgb = color
        else:
            r = p.add_run(); r.text = it
            r.font.name = SANS; r.font.size = Pt(size); r.font.color.rgb = color
    return s


def logo(sl, dark):
    sl.shapes.add_picture(LOGO_N if dark else LOGO_L,
                          Inches(W - M - 1.12), Inches(0.52), width=Inches(1.12))


def head(sl, dark, kick, title_parts, sub=None, bg=None, subw=None):
    """kicker + h2 + sub + rule. Devuelve la Y donde empieza el contenido."""
    base = NAVY if dark else (bg or ICE)
    rect(sl, 0, 0, W, H, base)
    logo(sl, dark)
    kc = CREAM if dark else BLUE
    text(sl, M, 0.62, CW - 1.4, 0.24, kick.upper(), SANS, 8, True, kc, spc=180)
    tc = SKY if dark else NAVY
    em = CREAM if dark else BLUE
    runs = [(t, {'color': em, 'italic': True} if it else {'color': tc})
            for t, it in title_parts]
    text(sl, M, 0.92, CW - 1.4, 0.62, [runs], SERIF, 29, False, tc, lead=1.02)
    y = 1.66
    if sub:
        sc = SKYM if dark else MUTED
        text(sl, M, y, subw or 8.9, 0.5, sub, SANS, 9.6, False, sc, lead=1.4)
        y += 0.52 if len(sub) < 150 else 0.72
    rl = rect(sl, M, y, CW, 0.012, SKYM if dark else LINE)
    return y + 0.2


def foot(sl, dark, left, num):
    c = SKYM if dark else MUTED
    text(sl, M, H - 0.62, 9.6, 0.22, left.upper(), SANS, 7, False, c, spc=120)
    text(sl, W - M - 0.9, H - 0.62, 0.9, 0.22, '%02d' % num, SANS, 7.5, True,
         CREAM if dark else BLUE, spc=120)


# variantes de tarjeta: (fill, line, label, head, body)
CARD = {
    'ice':  (ICE,   LINE, BLUE, NAVY, NAVY),
    'sky':  (SKY,   None, BLUE, NAVY, NAVY),
    'cream':(CREAM, None, BLUE, NAVY, NAVY),
    'blue': (BLUE,  None, CREAM, ICE,  ICE),
    'navy': (NAVY,  None, CREAM, SKY,  SKY),
    'ou':   (None,  OUTL, CREAM, SKY,  SKY),
}


def card(sl, kind, l, t, w, h, lbl=None, title=None, body=None, items=None,
         num=None, pad=0.26, tsize=13.5, bsize=9.2, center=False):
    f, ln, lc, hc, bc = CARD[kind]
    rect(sl, l, t, w, h, f, ln, radius=0.19)
    x = l + pad; iw = w - 2 * pad
    blocks = []
    if num:   blocks.append(('num', num))
    if lbl:   blocks.append(('lbl', lbl))
    if title: blocks.append(('ttl', title))
    if body:  blocks.append(('bdy', body if isinstance(body, list) else [body]))
    if items: blocks.append(('itm', items))
    # alturas estimadas
    hs = []
    for k, v in blocks:
        if k == 'num': hs.append(0.34)
        elif k == 'lbl': hs.append(0.22)
        elif k == 'ttl': hs.append(0.3 + 0.235 * tsize / 13.5 * max(1, len(v) / max(14, int(iw * 15))))
        elif k == 'bdy': hs.append(sum(0.17 + 0.145 * (len(p) // max(30, int(iw * 24))) for p in v) + 0.1 * (len(v) - 1))
        else: hs.append(sum(0.2 + 0.14 * (len(i[0] + (i[1] if isinstance(i, tuple) else '')) // max(28, int(iw * 22))) for i in v) if isinstance(v[0], tuple) else 0.22 * len(v))
    y = t + pad + (max(0, (h - 2 * pad - sum(hs)) / 2) if center else 0)
    for (k, v), bh in zip(blocks, hs):
        if k == 'num':
            text(sl, x, y, iw, 0.34, v, SERIF, 19, False, lc); y += 0.36
        elif k == 'lbl':
            text(sl, x, y, iw, 0.22, v.upper(), SANS, 7, True, lc, spc=130); y += 0.26
        elif k == 'ttl':
            text(sl, x, y, iw, bh, v, SERIF, tsize, False, hc, lead=1.05); y += bh + 0.08
        elif k == 'bdy':
            runs = []
            for p in v:
                if isinstance(p, tuple):
                    runs.append([(p[0], {'bold': True, 'color': bc}), (p[1], {'color': bc})])
                else:
                    runs.append([(p, {'color': bc})])
            text(sl, x, y, iw, bh, runs, SANS, bsize, False, bc, lead=1.4, space=5); y += bh + 0.06
        else:
            bullets(sl, x, y, iw, bh, v, bc, bsize - 0.2, lc); y += bh


def table(sl, l, t, w, rows, colw, dark=False, rowh=0.34, headh=0.3,
          size=8.4, hsize=7.2, rec_col=None, rec_rows=None):
    n, m = len(rows), len(rows[0])
    gt = sl.shapes.add_table(n, m, Inches(l), Inches(t), Inches(w),
                             Inches(headh + rowh * (n - 1))).table
    gt.first_row = False; gt.horz_banding = False
    tot = float(sum(colw))
    for c, cw in zip(gt.columns, colw):
        c.width = Emu(int(w * cw / tot * 914400))
    gt.rows[0].height = Emu(int(headh * 914400))
    for r in range(1, n):
        gt.rows[r].height = Emu(int(rowh * 914400))
    body_bg = NAVY if dark else ICE
    body_fg = SKY if dark else NAVY
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = gt.cell(ri, ci); cell.text = ''
            cell.margin_left = cell.margin_right = Emu(int(0.09 * 914400))
            cell.margin_top = cell.margin_bottom = Emu(int(0.025 * 914400))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = BLUE if (rec_col is not None and ci == rec_col) else NAVY
            elif rec_col is not None and ci == rec_col:
                cell.fill.fore_color.rgb = CREAM
            elif rec_rows and ri in rec_rows:
                cell.fill.fore_color.rgb = CREAM
            else:
                cell.fill.fore_color.rgb = body_bg
            p = cell.text_frame.paragraphs[0]
            p.line_spacing = 1.2
            r0 = p.add_run(); r0.text = val
            if ri == 0:
                r0.font.name = SANS; r0.font.size = Pt(hsize); r0.font.bold = True
                r0.font.color.rgb = ICE; _spc(r0, 110)
            else:
                rec = (rec_col is not None and ci == rec_col) or (rec_rows and ri in rec_rows)
                r0.font.name = SERIF if ci == 0 else SANS
                r0.font.size = Pt(size)
                r0.font.bold = (ci == 0)
                r0.font.color.rgb = NAVY if rec else body_fg
    return gt


def stat(sl, l, t, w, h, lbl, big, cap, dark=True, kind='ou'):
    f, ln, lc, hc, bc = CARD[kind]
    rect(sl, l, t, w, h, f, ln, radius=0.19)
    text(sl, l + 0.24, t + 0.2, w - 0.48, 0.2, lbl.upper(), SANS, 6.8, True, lc, spc=130)
    text(sl, l + 0.24, t + 0.44, w - 0.48, 0.52, big, SERIF, 25, False, hc, lead=1.0)
    text(sl, l + 0.24, t + h - 0.46, w - 0.48, 0.4, cap, SANS, 7.6, False, bc, lead=1.3)


def device(sl, l, t, w, h, l1, mid, dark=True, size=23):
    """El dispositivo ( ) Club: titular + escena entre corchetes + Club."""
    bg = NAVY if dark else SKY
    rect(sl, l, t, w, h, bg)
    fg = SKY if dark else NAVY
    ac = CREAM if dark else BLUE
    emb = EMB_N if dark else EMB_L
    text(sl, l + 0.5, t + h / 2 - 0.9, w - 1.0, 1.8,
         [[(l1 + u'\n', {'color': fg}),
           (u'( ' + mid + u' ) ', {'color': ac}),
           (u'Club', {'color': fg})]],
         SERIF, size, False, fg, lead=1.16)
    try:
        sl.shapes.add_picture(emb, Inches(l + 0.5), Inches(t + h - 1.0), width=Inches(0.62))
    except Exception:
        pass


def new(dark=False, bg=None):
    sl = prs.slides.add_slide(BLANK)
    return sl


# ================================================================ LAMINAS
# ---- 01 PORTADA
sl = new()
rect(sl, 0, 0, W, H, NAVY)
logo(sl, True)
text(sl, M, 0.62, 8.0, 0.24, u'DINERS CLUB INTERNATIONAL · EDICIONES GAMMA · OCTUBRE 2026',
     SANS, 8, True, CREAM, spc=180)
text(sl, M, 1.5, 7.4, 1.9, [[(u'Unlock Interesting\n', {'color': SKY}),
                             (u'llega a ', {'color': SKY}),
                             (u'Colombia', {'color': CREAM, 'italic': True})]],
     SERIF, 46, False, SKY, lead=1.0)
text(sl, M, 3.52, 6.6, 0.8,
     u'El lanzamiento de la marca Diners Club ante quien todavía no es Clubmember. '
     u'Doce semanas, con el ecosistema de Ediciones Gamma como medio de lanzamiento.',
     SANS, 10.4, False, SKYM, lead=1.45)
stats = [(u'Clubmembers en Colombia', u'169K', u'18.4K en segmento Black'),
         (u'Suscriptores de Revista Diners', u'146.000', u'del mismo perfil premium que el Clubmember'),
         (u'La revista con el nombre de la tarjeta', u'Desde 1963', u'una editorial con archivo propio')]
for i, (a, b, c) in enumerate(stats):
    stat(sl, M + i * 2.42, 4.68, 2.22, 1.5, a, b, c)
device(sl, 7.9, 1.5, 4.64, 4.68, u"It's the", u'Dining')
text(sl, M, H - 0.62, 6.0, 0.22, u'EDICIONES GAMMA · DOCUMENTO CONFIDENCIAL', SANS, 7, False, SKYM, spc=120)
text(sl, W - M - 3.0, H - 0.62, 3.0, 0.22, u'TRES MESES DESDE LA FIRMA', SANS, 7, True, CREAM, spc=120)

# ---- 02 EL ENCARGO
sl = new()
y = head(sl, False, u'El encargo', [(u'Esto es lo que el roadmap ', 0), (u'pidió', 1)],
         u'El encargo tiene dos objetivos: relanzamiento estratégico de marca y posicionamiento de beneficios. '
         u'Gamma lleva la marca a quien todavía no la tiene, en medios externos. Davivienda sigue hablando con los '
         u'Clubmembers actuales por sus propios canales. Esta propuesta cubre el primer trimestre.', subw=11.4)
card(sl, 'ice', M, y, 6.1, H - y - 0.78, lbl=u'Calendario del roadmap')
table(sl, M + 0.26, y + 0.56, 5.58,
      [[u'Momento', u'Qué ocurre'],
       [u'Q2 2026', u'DC Colombia empieza a adaptar las mejores prácticas'],
       [u'Q3 2026', u'Los socios de LATAM entregan su plan de medios'],
       [u'Q4 2026', u'Desarrollo de la campaña LATAM'],
       [u'Q1 2027', u'Los socios de LATAM lanzan la campaña de dining'],
       [u'2027', u'La app habilita ofertas locales cargadas por el emisor']],
      [1.0, 2.5], rowh=0.42, size=8.2)
card(sl, 'navy', M + 6.35, y, CW - 6.35, H - y - 0.78,
     lbl=u'Cómo se mapea sobre los cuatro elementos del programa',
     items=[(u'01 Dining Experiences. ', u'Experiencias OBI con cupo limitado'),
            (u'02 Rewards & Dining Offers. ', u'Mesa Diners Club, beneficio recurrente sobre aliados verificados'),
            (u'03 Chef & Dining Partnerships. ', u'Serie de video con cocineros y personajes locales'),
            (u'04 Dining Awards. ', u'Base editorial para un premio gastronómico colombiano')],
     body=[u'Ningún componente es invención local. Cada uno responde a un elemento ya definido por el programa global.'],
     center=True)
foot(sl, False, u'DCI Dining Program · Roadmap Colombia', 2)

# ---- 03 LA MARCA QUE SE LANZA
sl = new()
y = head(sl, True, u'La marca que se lanza',
         [(u'Unlock Interesting: una vida más ', 0), (u'interesante', 1)])
card(sl, 'ou', M, y, 5.5, 2.3, lbl=u'La verdad de marca',
     title=u'No es solo una tarjeta.\nEs un Club como ningún otro.',
     body=[u'El Enrichment Enthusiast no quiere comprar cosas ni experiencias. Quiere comprar una vida '
           u'más interesante, y Diners Club es su ruta para llegar ahí.'], tsize=17)
card(sl, 'ou', M, y + 2.45, 2.65, H - y - 3.23, lbl=u'A quién le habla',
     body=[u'Al Enrichment Enthusiast, que reinvierte tiempo, energía y dinero en su propia expansión '
           u'y la de su círculo cercano.'])
card(sl, 'ou', M + 2.85, y + 2.45, 2.65, H - y - 3.23, lbl=u'Contra qué compite',
     body=[u'Contra un lenguaje premium limitado a puntos y acceso, y contra lo premium producido en masa, '
           u'que hoy se siente poco interesante.'])
card(sl, 'navy', M + 5.75, y, CW - 5.75, H - y - 0.78,
     lbl=u'El sistema creativo, en tres reglas',
     items=[(u'Belonging. ', u"El titular abre con una invitación o una confirmación de pertenencia: We're, It's the, Welcome to the"),
            (u'Uniqueness. ', u'Los corchetes enmarcan la escena que representa la vida más interesante que Diners Club desbloquea'),
            (u'Distinctiveness. ', u'Todo titular termina en Club. Es lo que ninguna otra tarjeta puede decir')],
     body=[u'Se completa con el emblema del círculo partido, los dos colores primarios, rectángulos '
           u'redondeados y fotografía de escena real.'])
foot(sl, True, u'Guía de marca DCI · Brand Evolution, junio de 2026', 3)

# ---- 04 SU INVESTIGACION
sl = new()
y = head(sl, False, u'El punto de partida', [(u'Lo que su investigación ya ', 0), (u'señala', 1)])
card(sl, 'ice', M, y, 5.72, 2.52, lbl=u'Accesibilidad percibida, global, entre Clubmembers propios')
table(sl, M + 0.26, y + 0.56, 5.2,
      [[u'Atributo', u'Visa', u'Mastercard', u'Diners', u'Amex'],
       [u'Aceptada en mi zona', u'72', u'59', u'54', u'27'],
       [u'Aceptación en el mundo', u'76', u'56', u'28', u'24']],
      [2.1, .72, 1.0, .82, .72], rowh=0.42)
text(sl, M + 0.26, y + 1.92, 5.2, 0.5,
     u'“Resolver las barreras de aceptación es clave para retener a los clientes actuales, que de lo '
     u'contrario pueden optar por la competencia.”  Brand Health Tracker 2026 · 9.721 encuestados · 16 mercados',
     SANS, 7.4, False, MUTED, lead=1.35, italic=True)
card(sl, 'cream', M, y + 2.68, 5.72, H - y - 3.46, lbl=u'La consecuencia operativa',
     title=u'El plan no puede generar intención hacia comercios sin verificar',
     body=[u'Cada intento fallido en caja confirma la percepción que la investigación ya midió. '
           u'Por eso toda la comunicación de uso apunta únicamente a establecimientos con aceptación verificada.'],
     tsize=13)
card(sl, 'ice', M + 5.97, y, CW - 5.97, 2.52, lbl=u'Ofertas únicas de dining y deals relevantes, entre Clubmembers')
table(sl, M + 6.23, y + 0.56, 5.2,
      [[u'Marca', u'%'], [u'Diners Club', u'56'], [u'American Express', u'23'],
       [u'Visa', u'23'], [u'Mastercard', u'14']],
      [3.2, 1.0], rowh=0.33, rec_rows={1})
card(sl, 'navy', M + 5.97, y + 2.68, CW - 5.97, H - y - 3.46, lbl=u'Lo que esto define',
     title=u'Dining es el terreno donde la marca ya gana',
     body=[u'La ventaja en dining duplica a la del competidor más cercano. El plan no abre un frente '
           u'nuevo: concentra la inversión donde la marca ya tiene permiso.'], tsize=13)
foot(sl, False, u'Brand Health Tracker 2026 · Investigación propia de DCI', 4)

# ---- 05 LOS DOS PUBLICOS
sl = new()
y = head(sl, False, u'A quién llega este plan',
         [(u'Dos públicos, no uno. El encargo ', 0), (u'es el prospecto', 1)],
         u'Un beneficio de tarjeta solo alcanza a quien ya la tiene. Un medio editorial alcanza a quien todavía no.')
card(sl, 'ice', M, y, 5.9, H - y - 0.78, lbl=u'La brecha con prospectos, global')
table(sl, M + 0.26, y + 0.56, 5.38,
      [[u'Atributo', u'Diners', u'Visa', u'Mastercard'],
       [u'Entiendo del todo sus beneficios', u'7', u'47', u'42'],
       [u'Ofertas de dining relevantes', u'16', u'39', u'35'],
       [u'Aspiro a tenerla', u'11', u'42', u'38'],
       [u'Aceptada en mi zona', u'11', u'66', u'63']],
      [2.5, .85, .72, 1.1], rowh=0.44)
card(sl, 'cream', M + 6.15, y, CW - 6.15, (H - y - 0.78) / 2 - 0.08,
     lbl=u'Prospectos · 146.000 suscriptores', title=u'Comprensión y aspiración',
     body=[u'La revista entra a hogares del mismo perfil premium que todavía no son Clubmembers. Es el '
           u'único instrumento del plan que ataca la brecha de prospectos, y es el mandato de Gamma.'], tsize=13)
card(sl, 'navy', M + 6.15, y + (H - y - 0.78) / 2 + 0.08, CW - 6.15, (H - y - 0.78) / 2 - 0.08,
     lbl=u'Clubmembers · 169.000', title=u'Retención y uso',
     body=[u'Una razón concreta para elegir la tarjeta al pagar, en lugares con aceptación verificada. '
           u'Davivienda los comunica por sus canales y replica las piezas aprobadas.'], tsize=13)
foot(sl, False, u'El activo editorial es el que resuelve el segundo público', 5)

# ---- 06 EL ECOSISTEMA GAMMA
sl = new()
y = head(sl, True, u'El medio de lanzamiento',
         [(u'El ecosistema de Gamma lleva la marca a ', 0), (u'quien falta', 1)],
         u'Cada activo cumple un papel distinto en el lanzamiento. Todas las piezas llevan el sistema de marca de DCI.')
table(sl, M, y, 7.9,
      [[u'Activo Gamma', u'Qué aporta', u'Papel en el lanzamiento'],
       [u'Revista Diners, impreso', u'146.000 suscriptores del perfil premium', u'Irrupción y voz de marca'],
       [u'revistadiners.com.co', u'Profundidad y archivo', u'Artículo de origen, selección viva'],
       [u'Redes de Revista Diners', u'Video y cubrimiento', u'Manifiesto, serie con cocineros, reels'],
       [u'AXXIS, impreso y digital', u'Perfil 35+ con mirada de diseño', u'Cobertura ampliada, opción C'],
       [u'Libros', u'Unidad editorial en operación', u'Publicación anual con archivo y datos'],
       [u'Red de aliados verificados', u'Vitrinas en las mejores zonas', u'Sello, tent card y pieza de mesa']],
      [2.2, 2.5, 2.6], dark=True, rowh=0.52)
card(sl, 'cream', M + 8.15, y, CW - 8.15, H - y - 0.78, lbl=u'La regla de producción',
     title=u'Gamma produce.\nDCI aprueba.',
     body=[u'Cada pieza usa el dispositivo ( ) Club, los colores y la tipografía de la guía. '
           u'Nada sale sin visto bueno de marca.'], tsize=17, center=True)
foot(sl, True, u'Cifras de suscriptores: Revista Diners, base de circulación', 6)

# ---- 07 MAPA COMPETITIVO
sl = new()
y = head(sl, False, u'El espacio que queda libre', [(u'Dos ejes tomados, uno ', 0), (u'vacío', 1)],
         u'Los dos competidores de la categoría ocupan posiciones distintas y ninguno ocupa la tercera. '
         u'Esa es la razón de fondo de este posicionamiento.')
cols = [('ice', u'Visa', u'Everywhere you want to be',
         [(u'Promete cantidad. ', u'Aceptación universal, pago sin contacto, patrocinio de Mundial y Juegos Olímpicos, festivales masivos.'),
          u'Es la dimensión donde Diners pierde en su propia medición: 54 contra 72 en aceptación local. Disputarla no tiene salida.']),
        ('ice', u'Mastercard', u'Priceless',
         [(u'Promete intensidad. ', u'Momentos memorables, acceso exclusivo, patrocinio deportivo y mesas con chefs premiados.'),
          u'Tiene plataforma global de reserva y presupuesto de patrocinio. Competir de frente es hacer lo mismo, más pequeño.']),
        ('cream', u'Diners Club', u'Saber dónde',
         [(u'Promete criterio. ', u'Visa dice dónde se puede pagar. Mastercard dice qué momento comprar. Ninguna dice cuáles valen la pena.'),
          u'El criterio no se compra con presupuesto de medios. Se acumula con años de cubrimiento.'])]
ch = H - y - 1.72
for i, (k, lbl, ttl, bd) in enumerate(cols):
    card(sl, k, M + i * 3.98, y, 3.73, ch, lbl=lbl, title=ttl, body=bd, tsize=14)
card(sl, 'navy', M, y + ch + 0.16, CW, 0.78, lbl=u'Y un público desatendido',
     body=[u'Mastercard declara como objetivo los segmentos premium, la Gen Z y los apasionados del deporte. '
           u'Visa apunta al usuario masivo y adoptante digital. El Clubmember clásico de 45 a 55 años no es '
           u'el foco de ninguna de las dos.'])
foot(sl, False, u'Mapa de posicionamiento de la categoría en Colombia', 7)

# ---- 08 LAS TRES CAPAS
sl = new()
y = head(sl, False, u'La arquitectura', [(u'Dos objetivos, tres ', 0), (u'capas', 1)],
         u'El relanzamiento de marca se sostiene en la capa editorial, que alcanza a quien todavía no es '
         u'Clubmember. El posicionamiento de beneficios se sostiene en el beneficio recurrente y en las '
         u'experiencias por invitación. La aceptación verificada atraviesa los dos.', subw=11.4)
capas = [(u'01', u'El beneficio', u'Mesa Diners Club',
          u'Beneficio gastronómico recurrente con día fijo, sobre una selección de aliados con aceptación verificada.',
          [u'Día fijo semanal', u'Datáfono verificado como requisito', u'Bogotá en el año uno']),
         (u'02', u'El descubrimiento', u'La Mesa Diners',
          u'La sección editorial donde vive el beneficio. Convierte el descuento en criterio y alimenta el catálogo que la app cargará en 2027.',
          [u'Sección fija mensual en la revista', u'Sección permanente en web', u'Serie de videos']),
         (u'03', u'El estatus', u'Experiencias OBI',
          u'Only By Invitation. Acceso limitado a experiencias que no se compran, dirigido al segmento Black.',
          [u'Seis experiencias en el año', u'Cupo corto, por invitación', u'Cubrimiento editorial incluido'])]
for i, (n_, lbl, ttl, bd, it) in enumerate(capas):
    card(sl, 'navy', M + i * 3.98, y, 3.73, H - y - 0.78, num=n_, lbl=lbl, title=ttl,
         body=[bd], items=it, tsize=15)
foot(sl, False, u'Las tres operan sobre la misma red de aliados verificados', 8)

# ---- 09 MESA DINERS CLUB
sl = new()
y = head(sl, False, u'Capa 1 · El beneficio', [(u'Mesa ', 0), (u'Diners Club', 1)],
         u'Adaptación colombiana de la mecánica de descuento en días de baja reserva, que el programa '
         u'global reconoce como práctica probada.', subw=7.6)
card(sl, 'ice', M, y, 4.1, H - y - 0.78, lbl=u'Cómo opera', title=u'Un día fijo, una lista corta',
     items=[(u'Beneficio: ', u'porcentaje sobre la cuenta'),
            (u'Frecuencia: ', u'día fijo semanal, para construir hábito'),
            (u'Cobertura: ', u'Bogotá en el año uno'),
            (u'Requisito del aliado: ', u'Diners verificada en el datáfono'),
            (u'Condición: ', u'reserva obligatoria')], tsize=14)
card(sl, 'cream', M + 4.35, y, 4.1, H - y - 0.78, lbl=u'Qué lo diferencia del modelo de descuento',
     title=u'El aliado paga la visibilidad',
     body=[u'En lugar de financiar un descuento profundo, el establecimiento paga por aparecer en la '
           u'selección editorial. El costo neto del programa baja y la oferta crece sin inversión adicional.',
           u'Es una oferta menos agresiva y una economía que mejora con la escala. Quedamos abiertos a '
           u'modelar la alternativa de descuento financiado.'], tsize=14)
device(sl, M + 8.7, y, CW - 8.7, H - y - 0.78, u'A fixed day.', u'A short list', size=17)
foot(sl, False, u'Mapea sobre Rewards & Dining Offers', 9)

# ---- 10 LA MESA DINERS
sl = new()
y = head(sl, True, u'Capa 2 · El descubrimiento',
         [(u'Un descuento no se recuerda, un criterio ', 0), (u'sí', 1)],
         u'La selección de los lugares, cocineros y rituales que definen el buen comer en Bogotá, '
         u'con el estándar editorial de la Revista Diners.')
sop = [(u'Soporte 1', u'Sección fija mensual',
        u'Sección nueva, complementaria al programa. No ocupa espacio existente: se crea para esto, en posición fija.'),
       (u'Soporte 2', u'Sección permanente en web',
        u'La selección viva y buscable, enlazada con la lista de aliados verificados.'),
       (u'Soporte 3', u'Serie de video',
        u'Los protagonistas en video editorial, distribuido en canales propios y en los del invitado.')]
sh_ = H - y - 1.98
for i, (lbl, ttl, bd) in enumerate(sop):
    card(sl, 'ou', M + i * 3.98, y, 3.73, sh_, lbl=lbl, title=ttl, body=[bd], tsize=14)
card(sl, 'cream', M, y + sh_ + 0.18, CW, 1.02, lbl=u'La regla que protege el criterio',
     title=u'La entrada la define el editor',
     body=[u'Pagar amplía la ficha, no compra el lugar en la selección. Sin esa regla la sección se lee '
           u'como publirreportaje y se pierde la autoridad editorial, que es lo único que ninguna agencia puede replicar.'],
     tsize=13)
foot(sl, True, u'Mapea sobre Bespoke Dining Content', 10)

# ---- 11 OBI
sl = new()
y = head(sl, False, u'Capa 3 · El estatus', [(u'Experiencias ', 0), (u'Only By Invitation', 1)],
         u'Seis experiencias gastronómicas en el año, con cupo limitado y acceso por invitación.', subw=7.6)
obi = [('cream', u'La escasez es el punto', u'18.400 Black · 150 cupos',
        u'El cupo corto es el beneficio, no una limitación de logística. Lo que se reparte a todos deja de distinguir a alguien.'),
       ('ice', u'Rendimiento triple', u'Se produce una vez, se distribuye tres',
        u'Evento, reel de cubrimiento y artículo permanente. El contenido alcanza a quien no estuvo, incluidos los prospectos.'),
       ('ice', u'Prueba social', u'Testimonios de Clubmembers',
        u'El propio tracker recomienda usar testimonios de Clubmembers para mostrar relevancia a los prospectos. Cada OBI los genera.')]
for i, (k, lbl, ttl, bd) in enumerate(obi):
    card(sl, k, M, y + i * ((H - y - 0.78) / 3), 7.5, (H - y - 0.78) / 3 - 0.16,
         lbl=lbl, title=ttl, body=[bd], tsize=13.5)
device(sl, M + 7.75, y, CW - 7.75, H - y - 0.78, u'What everyone gets stops setting', u'anyone', size=16)
foot(sl, False, u'Mapea sobre Dining Experiences', 11)

# ---- 12 EL SELLO
sl = new()
y = head(sl, False, u'La pieza que resuelve la aceptación', [(u'Un sello en la ', 0), (u'vitrina', 1)],
         u'Material de punto de venta en la entrada de cada aliado verificado, con el formato que la '
         u'franquicia ya usa en la región. Es la pieza de menor costo unitario del plan y la de mayor rendimiento.',
         subw=8.4)
res = [(u'Resuelve 1', u'Señala dónde sí funciona',
        u'El Clubmember ve el sello antes de entrar. Es la respuesta física a la barrera de aceptación que el tracker identifica como riesgo de retención.'),
       (u'Resuelve 2', u'Convierte la selección en estatus',
        u'Estar en la lista prestigia al restaurante. No es publicidad que se tolera: es un reconocimiento que se exhibe.'),
       (u'Resuelve 3', u'Sostiene la economía de la red',
        u'Es el argumento de venta para que el aliado entre y pague su visibilidad. Lo que compra es pertenecer a una selección firmada por la Revista Diners.'),
       (u'Resuelve 4', u'Entre 20 y 40 vitrinas sin costo de pauta',
        u'En las mejores zonas de Bogotá, visibles todo el año. Con código para llevar al catálogo digital y, desde 2027, a la app.')]
gh = (H - y - 0.78) / 2 - 0.08
for i, (lbl, ttl, bd) in enumerate(res):
    card(sl, 'cream' if i < 2 else 'ice', M + (i % 2) * 4.3, y + (i // 2) * (gh + 0.16), 4.05, gh,
         lbl=lbl, title=ttl, body=[bd], tsize=13)
device(sl, M + 8.7, y, CW - 8.7, H - y - 0.78, u'Acceptance is solved at the', u'door', size=16)
foot(sl, False, u'Mapea sobre Dining Merchant Acceptance Push', 12)

# ---- 13 BENCHMARK DE CATEGORIA
sl = new()
y = head(sl, False, u'Benchmark de categoría · Colombia',
         [(u'La competencia está convirtiendo los beneficios\nen ', 0), (u'experiencias', 1), (u'.', 0)],
         u'Las tarjetas premium ya no compiten solo por descuentos o aceptación. Compiten por acceso, '
         u'momentos y espacios que hacen visible la pertenencia.', subw=9.6)
y += 0.34
text(sl, M, y - 0.3, 6.0, 0.22, u'¿QUÉ ESTÁ PASANDO EN COLOMBIA?', SANS, 7.4, True, SAND, spc=140)
comp = [(u'Mastercard · Priceless', u'Acceso que no se compra',
         u'Leo con Leonor Espinosa · Prudencia con Mario Rosero · experiencias de música, '
         u'gastronomía y cultura · nueva plataforma World Legend.',
         u'EXPERIENCIAS PRIVADAS + TALENTO + RELATO'),
        (u'American Express', u'La marca se vive',
         u'Lounge Bar en Movistar Arena · Lounge in Lounge en El Dorado · Dining Program en varias ciudades.',
         u'ESPACIOS FÍSICOS + HOSPITALIDAD'),
        (u'Visa', u'Grandes momentos',
         u'FIFA · Juegos Olímpicos · plataformas deportivas y culturales. En Colombia, conexión '
         u'con grandes momentos de consumo y turismo.',
         u'PATROCINIO + ESCALA CULTURAL'),
        (u'DAVIbank / One Rewards', u'Beneficio que incentiva uso',
         u'Salas VIP · traslados al aeropuerto · beneficios de viaje y condiciones asociadas al uso de la tarjeta.',
         u'EL BENEFICIO GENERA COMPORTAMIENTO'),
        (u'Bancolombia', u'Arquitectura premium',
         u'World Legend Mastercard · Icon American Express · salas VIP · gastronomía · entretenimiento · viajes.',
         u'EL ECOSISTEMA PREMIUM SE VUELVE ESTÁNDAR'),
        (u'Nu Colombia', u'Marca dentro de la cultura',
         u'Activación alrededor de Inter Miami × Atlético Nacional · arte colombiano en el kit de '
         u'bienvenida · identidad visual fuerte.',
         u'LA MARCA SALE DEL PRODUCTO Y ENTRA EN LA CULTURA')]
cw_ = (CW - 2 * 0.26) / 3
gh = (H - y - 0.86) / 2 - 0.14
for i, (brand, ttl, bd, arrow) in enumerate(comp):
    cl = M + (i % 3) * (cw_ + 0.26)
    ct = y + (i // 3) * (gh + 0.28)
    rect(sl, cl, ct, cw_, gh, CREAM, SAND, radius=0.19)
    text(sl, cl + 0.26, ct + 0.2, cw_ - 0.52, 0.2, brand.upper(), SANS, 7, True, SAND, spc=130)
    text(sl, cl + 0.26, ct + 0.46, cw_ - 0.52, 0.34, ttl, SERIF, 14.5, True, NAVY, lead=1.05)
    text(sl, cl + 0.26, ct + 0.88, cw_ - 0.52, gh - 1.38, bd, SANS, 8.4, False, NAVY, lead=1.4)
    text(sl, cl + 0.26, ct + gh - 0.42, cw_ - 0.52, 0.32,
         u'→  ' + arrow, SANS, 7, True, SAND, spc=110, lead=1.2)
foot(sl, False, u'Fuentes: plataformas oficiales de Mastercard Priceless, American Express Colombia, '
                u'Visa, DAVIbank, Bancolombia y Nu Colombia · revisión 2026', 13)

# ---- 14 EL JUEGO ESTA CAMBIANDO
sl = new()
y = head(sl, False, u'La lectura estratégica', [(u'El juego está ', 0), (u'cambiando', 1), (u'.', 0)],
         u'La categoría está pasando de entregar beneficios a construir razones para pertenecer.')
esc = [(u'01', u'DESCUENTO', u'“Ahorra con tu tarjeta”'),
       (u'02', u'BENEFICIO', u'“Obtén algo extra”'),
       (u'03', u'ACCESO', u'“Entra donde otros no”'),
       (u'04', u'EXPERIENCIA', u'“Vive algo que recordarás”'),
       (u'05', u'PERTENENCIA', u'“Esto significa algo de ti”')]
cw_ = (CW - 4 * 0.24) / 5
et = y + 1.0
eh = 1.95
for i, (n_, ttl, q) in enumerate(esc):
    last = (i == 4)
    cl = M + i * (cw_ + 0.24)
    rect(sl, cl, et, cw_, eh, NAVY if last else CREAM, None if last else SAND, radius=0.19)
    nc = SAND if last else SAND
    text(sl, cl + 0.26, et + 0.26, cw_ - 0.52, 0.2, n_, SANS, 7.4, True, nc, spc=130)
    text(sl, cl + 0.26, et + 0.56, cw_ - 0.52, 0.4, ttl, SERIF, 15.5, True,
         WHITE if last else NAVY, spc=40)
    text(sl, cl + 0.26, et + 1.06, cw_ - 0.52, 0.5, q, SERIF, 9.5, False,
         SKY if last else MUTED, lead=1.3, italic=True)
foot(sl, False, u'Benchmark de activaciones y plataformas documentadas en Colombia · 2025-2026', 14)

# ---- 15 EL CICLO
sl = new()
y = head(sl, False, u'Cómo cierra el ciclo', [(u'De la revista al dato, y del dato a la ', 0), (u'app', 1)],
         u'Cada pieza del plan alimenta la siguiente. El año uno construye la base de datos propia que el '
         u'ecosistema digital de la franquicia podrá explotar desde 2027.', subw=11.4)
ciclo = [(u'01', u'Alcance', u'La revista entra al hogar',
          u'Doce veces al año, a 146.000 hogares del perfil premium. Clubmembers y prospectos.'),
         (u'02', u'Captura', u'Código de activación',
          u'Cada ejemplar lleva un código único. Convierte alcance anónimo en registro rastreable por cohorte.'),
         (u'03', u'Uso', u'Redención verificada',
          u'El beneficio se redime en aliados con aceptación confirmada. Genera dato de categoría, zona y frecuencia.'),
         (u'04', u'Escala · desde 2027', u'La app carga la selección',
          u'El hub habilita que el emisor cargue sus ofertas locales por ubicación. La lista de aliados verificados es ese contenido.')]
cw_ = (CW - 3 * 0.24) / 4
for i, (n_, lbl, ttl, bd) in enumerate(ciclo):
    card(sl, 'navy' if i < 3 else 'cream', M + i * (cw_ + 0.24), y, cw_, H - y - 1.78,
         num=n_, lbl=lbl, title=ttl, body=[bd], tsize=13)
card(sl, 'ice', M, H - 1.58, CW, 0.8, lbl=u'Por qué importa el orden',
     body=[u'La app necesita contenido curado y aliados verificados para ser útil el día que se encienda. '
           u'El año uno produce las dos cosas. Sin esa base, la funcionalidad llega vacía.'])
foot(sl, False, u'Coincide con el CHS Benefits Hub', 15)

# ---- 16 MULTIPLICADORES
sl = new()
y = head(sl, False, u'Eficiencia del modelo', [(u'Multiplicadores de la ', 0), (u'inversión', 1)],
         u'Seis mecanismos que hacen que una misma inversión produzca resultado en más de un frente. '
         u'Están incorporados en las tres opciones.', subw=11.4)
mult = [(u'01', u'Rendimiento por pieza', u'Una experiencia, cinco salidas',
         u'Un costo de producción y cinco superficies: evento, reel, artículo, La Mesa Diners y newsletter.'),
        (u'02', u'Traslado de costo', u'El aliado produce la experiencia',
         u'Aporta cocina, espacio y equipo a cambio del cubrimiento. La línea de mayor costo se reduce a coordinación.'),
        (u'03', u'Activo existente', u'Archivo editorial propio',
         u'El archivo de la revista y los datos del programa se convierten en una publicación anual.'),
        (u'04', u'Red como canal', u'Cuentas de aliados publicando',
         u'Cada aliado promueve el beneficio porque le genera tráfico de alto ticket. Distribución sin inversión en medios.'),
        (u'05', u'Ajuste sin descuento', u'AXXIS como variable',
         u'Ficha técnica idéntica con tarifa menor. Amplía cobertura sin afectar el valor del plan.'),
        (u'06', u'Contenido del socio', u'La selección la nominan ellos',
         u'Contenido sin costo y datos de preferencia, con los Clubmembers nominando lugares.')]
gh = (H - y - 0.78) / 2 - 0.1
cw_ = (CW - 2 * 0.24) / 3
for i, (n_, lbl, ttl, bd) in enumerate(mult):
    card(sl, 'navy' if i < 3 else 'cream', M + (i % 3) * (cw_ + 0.24), y + (i // 3) * (gh + 0.2),
         cw_, gh, num=n_, lbl=lbl, title=ttl, body=[bd], tsize=13)
foot(sl, False, u'Las seis operan en las tres opciones', 16)

# ---- 17 DOCE SEMANAS
sl = new()
y = head(sl, True, u'El plan', [(u'Doce semanas, cuatro ', 0), (u'bloques', 1)],
         u'Arrancan con la firma, sin depender del mes del calendario. Línea de campaña propuesta: '
         u'Saber dónde, sujeta a aprobación de DCI.', subw=11.4)
blo = [(u'Semanas 1 a 3', u'Instalación', u'Que la marca vuelva a existir en la conversación.',
        u'Se mide: alcance sobre perfil premium y recordación asistida'),
       (u'Semanas 4 a 6', u'Demostración', u'Que el criterio se vea, no se declare.',
        u'Se mide: tiempo de lectura de la sección y reproducciones completas de video'),
       (u'Semanas 7 a 9', u'Prueba', u'Que haya evidencia de terceros.',
        u'Se mide: menciones ganadas y asistencia a la experiencia sobre invitados'),
       (u'Semanas 10 a 12', u'Conversión', u'Que haya una ruta clara a pedir la tarjeta.',
        u'Se mide: solicitudes de apertura digital atribuibles a nuestros medios')]
cw_ = (CW - 3 * 0.24) / 4
bh_ = H - y - 1.96
for i, (lbl, ttl, bd, med) in enumerate(blo):
    card(sl, 'ou', M + i * (cw_ + 0.24), y, cw_, bh_, lbl=lbl, title=ttl, body=[bd, med], tsize=15)
qc = [u'El posicionamiento instalado y medido: el año entra con línea base de recordación.',
      u'La red de aliados verificados, que es el catálogo que el programa de beneficios necesita.',
      u'El primer dato propio de conversión, para dimensionar el año con un costo por solicitud real.']
for i, q in enumerate(qc):
    card(sl, 'cream', M + i * (3.98), y + bh_ + 0.18, 3.73, 0.98, lbl=u'Queda construido', body=[q], bsize=8.4)
foot(sl, True, u'La decisión del año se toma con dato propio', 17)

# ---- 18 TRES ALCANCES
sl = new()
y = head(sl, False, u'Las opciones', [(u'Tres alcances para el ', 0), (u'trimestre', 1)],
         u'La misma estrategia, en tres alcances. Las tres corren doce semanas desde la firma y las tres '
         u'dejan medición de conversión a solicitud digital.', subw=11.4)
ops = [('ice', u'Opción A', u'Instalación',
        u'Solo el primer objetivo. Instalar la marca ante quien no la tiene, sin montar todavía la red de aliados.',
        [u'Pauta dinámica los 3 meses', u'Una pieza de alto impacto de apertura',
         u'Video manifiesto y su amplificación', u'Sección editorial en la revista', u'Solo Revista Diners']),
       ('cream', u'Opción B · Recomendado', u'Las doce semanas',
        u'Los dos objetivos operando y el primer dato propio de conversión.',
        [u'Todo lo de la opción A', u'20 aliados con aceptación verificada', u'Sello en punto de venta con código',
         u'Una experiencia por invitación', u'Pauta digital en 4 canales', u'Serie de video con cocineros',
         u'Tablero de conversión']),
       ('ice', u'Opción C', u'Ampliada',
        u'Las doce semanas con mayor cobertura de audiencia y de red.',
        [u'Todo lo de la opción B', u'AXXIS sumado al ecosistema', u'40 aliados verificados',
         u'Dos experiencias por invitación', u'Línea de video con personaje invitado', u'LinkedIn y programática'])]
for i, (k, lbl, ttl, bd, it) in enumerate(ops):
    card(sl, k, M + i * 3.98, y, 3.73, H - y - 0.78, lbl=lbl, title=ttl, body=[bd], items=it,
         tsize=15, bsize=8.6)
foot(sl, False, u'Tarifas 2026 sujetas a IVA', 18)

# ---- 19 COMPARATIVO
sl = new()
y = head(sl, False, u'Comparativo', [(u'Qué incluye cada ', 0), (u'opción', 1)])
table(sl, M, y, CW,
      [[u'Componente', u'A · Esencial', u'B · Completo', u'C · Insignia'],
       [u'Duración', u'12 semanas', u'12 semanas', u'12 semanas'],
       [u'Marcas', u'Solo Diners', u'Solo Diners', u'Diners + AXXIS'],
       [u'Alto impacto impreso', u'1 pieza', u'1 pieza', u'2 piezas'],
       [u'Pauta digital', u'No incluye', u'4 canales', u'6 canales'],
       [u'Aliados verificados', u'No incluye', u'20', u'40'],
       [u'Sello en punto de venta', u'No incluye', u'Sí', u'Sí'],
       [u'Experiencias por invitación', u'No incluye', u'1', u'2'],
       [u'Video', u'Manifiesto', u'Manifiesto + cocineros', u'+ personaje invitado'],
       [u'Medición de solicitudes', u'Parcial', u'Completa', u'Completa']],
      [2.7, 2.1, 2.4, 2.2], rowh=0.37, rec_col=2)
card(sl, 'cream', M, H - 1.56, CW, 0.78, lbl=u'Por qué recomendamos la opción B',
     body=[u'Es el alcance mínimo que activa los dos objetivos y deja medición de punta a punta. '
           u'La opción A instala marca sin construir la red. La C amplía cobertura cuando el presupuesto lo permita.'])
foot(sl, False, u'Tarifas 2026 sujetas a IVA', 19)

# ---- 20 ECUADOR
sl = new()
y = head(sl, True, u'Caso de referencia en la región', [(u'Lo que ya funcionó en ', 0), (u'Ecuador', 1)],
         u'La franquicia construyó su territorio gastronómico apoyada en su propia casa editorial.')
ec = [(u'Recordación de marca', u'76%', u'en socios, en experiencias gastronómicas'),
      (u'Incremento en socios', u'+14%', u'en el territorio'),
      (u'Recetas compiladas', u'1.053', u'en libros producidos por su editorial')]
for i, (a, b, c) in enumerate(ec):
    stat(sl, M + i * 2.52, y, 2.32, 1.5, a, b, c)
card(sl, 'cream', M, y + 1.66, 7.3, H - y - 2.44, lbl=u'Lo que demuestra',
     title=u'La casa editorial es el motor',
     body=[u'El territorio no se construyó solo con pauta: se construyó con contenido propio, publicaciones '
           u'y series, producidas por la editorial del emisor.',
           u'Ediciones Gamma ocupa en Colombia la misma posición, con una revista homónima desde 1963, '
           u'146.000 suscriptores y una unidad de Libros en operación.'], tsize=15)
device(sl, M + 7.55, y, CW - 7.55, H - y - 0.78, u'The territory was built by an', u'editorial house', size=15)
foot(sl, True, u'Carolina Ramirez · Ediciones Gamma · Diners Club del Ecuador', 20)

prs.save(OUT)
print('OK ->', OUT, len(prs.slides._sldIdLst), 'laminas')
